"""Template-derived anonymous competition drafts; no performance inventions."""
from pathlib import Path
import json,re,zipfile,io,hashlib
from lxml import etree
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1];REF=ROOT/'materials/references';OUT=ROOT/'materials/final';OUT.mkdir(exist_ok=True)
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';NS={'w':W};counts={}
TITLE='能智核——公共建筑风阀疑点核验与补证工作台'
def nchar(s):return len(re.sub(r'\s+','',s))
def font(run,size=12,bold=False,heading=False):
 run.font.name='SimHei' if heading else 'SimSun';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 rpr=run._element.get_or_add_rPr();rf=rpr.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
 for k in list(rf.attrib):rf.attrib.pop(k)
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimHei' if heading else 'SimSun')
def pformat(p,size=12,heading=False):
 pf=p.paragraph_format;pf.space_before=Pt(5 if heading else 0);pf.space_after=Pt(6 if heading else 4);pf.line_spacing=1.22
 pf.keep_with_next=heading;pf.keep_together=heading;pf.widow_control=True;pf.first_line_indent=None
 for x in ['snapToGrid','pBdr','pageBreakBefore']:
  for el in p._p.xpath('w:pPr/w:'+x):el.getparent().remove(el)
 for r in p.runs:font(r,size,heading,heading)
def setup(i):
 d=Document(next(REF.glob(f'附件{i}*.docx')))
 if 'Title' not in d.styles:
  from docx.enum.style import WD_STYLE_TYPE
  d.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
 for style in d.styles:
  if style.type in (1,2):
   style.font.name='SimSun';style.font.size=Pt(12)
   if style.element.rPr is not None:
    for rf in style.element.rPr.findall(qn('w:rFonts')):
     for k in list(rf.attrib):rf.attrib.pop(k)
     for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimSun')
 for sec in d.sections:
  for grid in sec._sectPr.findall(qn('w:docGrid')):sec._sectPr.remove(grid)
  for area in [sec.header,sec.footer]:
   for p in area.paragraphs:pformat(p,9)
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=TITLE;d.core_properties.subject='智慧能源与环境';d.core_properties.comments=''
 return d
def para(d,text,size=12,heading=False):
 p=d.add_paragraph(text);pformat(p,size,heading)
 if size<=9.5:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 return p
def h(d,text):return para(d,text,14,True)
def sub(d,text):return para(d,text,12,True)
def newbody(i):
 d=setup(i)
 for el in list(d._element.body):
  if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
 para(d,f'附件{i}：',10.5)
 p=para(d,'第十二届中国研究生智慧城市技术与创意设计大赛',14,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,'创意设计赛项目'+('说明书' if i==2 else '商业计划书'),16,True);p.style=d.styles['Title'];pformat(p,16,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,TITLE,12,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,'参赛方向：智慧能源与环境',10.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 return d
def page(d):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
def picture(d,path,caption):
 p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(path),width=Inches(5.70))
 p=para(d,caption,9.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(d,headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 if widths:
  for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,x in enumerate(headers):t.rows[0].cells[j].text=x
 for row in rows:
  c=t.add_row().cells
  for j,x in enumerate(row):c[j].text=str(x)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for tag in ['top','left','bottom','right']:
    line=OxmlElement('w:'+tag);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'BBCBC5');b.append(line)
   tcpr.append(b)
   if i==0:s=OxmlElement('w:shd');s.set(qn('w:fill'),'EDF3EF');tcpr.append(s)
   for p in c.paragraphs:
    pformat(p,10.5,i==0);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 return t
def sanitize_save(d,path):
 buf=io.BytesIO();d.save(buf);zin=zipfile.ZipFile(buf);entries={n:zin.read(n) for n in zin.namelist() if not n.startswith('word/fonts/')}
 fontxml=etree.Element(qn('w:fonts'),nsmap={'w':W})
 for name in ['SimSun','SimHei']:
  f=etree.SubElement(fontxml,qn('w:font'));f.set(qn('w:name'),name)
 entries['word/fontTable.xml']=etree.tostring(fontxml,xml_declaration=True,encoding='UTF-8',standalone=True)
 entries.pop('word/_rels/fontTable.xml.rels',None)
 entries.pop('docProps/custom.xml',None)
 for name,raw in list(entries.items()):
  if not name.endswith(('.xml','.rels')):continue
  root=etree.fromstring(raw)
  if name=='[Content_Types].xml':
   for x in list(root):
    if x.get('Extension')=='odttf' or x.get('PartName','').startswith('/word/fonts/') or x.get('PartName')=='/docProps/custom.xml':root.remove(x)
  if name=='_rels/.rels':
   for x in list(root):
    if x.get('Type','').endswith('/custom-properties'):root.remove(x)
  if name=='docProps/app.xml':
   for x in list(root):
    local=etree.QName(x).localname
    if local=='Application':x.text='Office Open XML'
    elif local=='Pages':x.text=str(5 if '附件2' in path.name else 3)
    elif local in ['Words','Characters','CharactersWithSpaces','Lines','Paragraphs','Company','Manager','Template']:root.remove(x)
  # Strip subset fonts, theme fallback and unrequested template personal metadata.
  if name.startswith('word/'):
   for rf in root.findall('.//'+qn('w:rFonts')):
    ishead=rf.get(qn('w:ascii'))=='SimHei'
    for k in list(rf.attrib):rf.attrib.pop(k)
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimHei' if ishead else 'SimSun')
   for el in root.findall('.//'+qn('w:embedRegular'))+root.findall('.//'+qn('w:embedBold'))+root.findall('.//'+qn('w:embedItalic'))+root.findall('.//'+qn('w:embedBoldItalic')):el.getparent().remove(el)
  entries[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in entries.items():z.writestr(n,b)
def brief():
 d=setup(1);t=d.tables[0]
 for i,p in enumerate(d.paragraphs):
  if p.text.strip()=='(模板)':p.text=TITLE
  pformat(p,14 if i in [1,2] else 10.5,i in [1,2])
  if i in [1,2,3]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for row in t.rows:
  row.height=None
  for pr in row._tr.xpath('w:trPr/w:trHeight'):pr.getparent().remove(pr)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 def cell(r,col,text):t.cell(r,col).text=text
 cell(0,1,TITLE)
 tracks=['低空经济与空天科技','智慧交通与数字孪生','智能建造与城市更新','未来城市与太空人居','智慧文旅与乡村振兴','智慧能源与环境','韧性城市与智慧应急','AI与具身智能','智慧生活与大健康','城市运维与安全','智能材料与绿色建造','自由赛道','企业揭榜赛道','地方特色赛道']
 options=[f'{i+1}. {s}'+(' ☑' if i==5 else ' □') for i,s in enumerate(tracks)]
 cell(2,1,'1—11为定向赛道：\n'+'\n'.join('    '.join(options[i:i+2]) for i in range(0,14,2)))
 texts={6:'公共建筑用能告警发生后，管理人员仍需核对数据是否可信、设备是否处于可检验工况，以及下一步应补哪个点位或时段。能智核面向这一核查环节，提供数据质量与设备证据分开呈现的工作台。当前以公开单风道空调机组仿真数据实现核验，以BDG2独立展示真实用能回放，服务于有边界、可复查的异常核查。',
 8:'把“告警之后先查什么”转化为可操作流程：先校验单位和真实时间间隔，再检查每项故障假设的适用条件；有支持证据时给出候选，缺点位、缺工况或证据冲突时给出具体补证要求。已有本地软件、三个案例、公开原始CSV导入和可阅读核查卡；真人价值对照包待独立复核与招募。特色在证据约束的产品流程，不主张已证明新的诊断算法。',
 10:'用户选择LBNL单风道机组公开仿真日片段，映射时间列、点位、单位及采样间隔。程序显式转换温度与指令单位，将完整1分钟桶聚合为5分钟；缺失值保留，间隔不符则拒绝。服务、诊断证据、质量卡和导出复用同一质量对象。通过质量校验后，程序按启停、稳定运行、温差和关阀持续时长生成证据窗口，分别核验风路、盘管和运行工况。界面呈现支持与限制证据、分支状态及补证条件；修正数据或修改证据后真实重算。导出的HTML核查卡可离线阅读和打印，并记录源数据哈希、映射及运行编号。设备范围限于当前仿真配置，真实建筑回放不参与故障定位。',
 12:'潜在使用方为公共建筑运维团队及能源服务方，拟提供按设备点位配置的本地核查工具和接入服务。现阶段仅完成公开数据原型与功能验证，尚无试点、合作或收入事实。后续需经合法数据接入、工程师复核和现场验证后再确定收费。节能机会筛查为附属方向，当前不承诺节能量、人工工时或运维收益。'}
 for r,s in texts.items():cell(r,1,s)
 counts['简表']={str(r):{'chars':nchar(s),'cap':600 if r==10 else 300} for r,s in texts.items()}
 cell(13,2,'创意设计类；基础研究类；\n☑软硬件开发类；工程实施类')
 cell(14,2,'国际领先；国际先进；国内领先；国内先进；\n国内一般；☑无法判断')
 cell(15,2,'□硬件 ☑软件 □工艺 □方法 □服务 □商业模式\n□其他')
 cell(16,2,'□第1级 □第2级 ☑第3级 □第4级\n□第5级 □第6级 □第7级 □第8级 □第9级\n自评：已完成公开仿真概念验证，尚无现场应用验证。')
 cell(17,2,'国际产品/技术/标准；国内产品/技术/标准：未开展独立对标。\n固定阈值为自建工程基线，非行业产品；没有优势证据。')
 cell(19,2,'☑质量提升：消除两个已复现的质量/时长错误。\n成本降低、效率提升、交付周期：未验证。\n☑新功能实现：原始CSV导入至可阅读核查卡。')
 nested=t.cell(18,2).tables[0]
 vals=[['指标名称','指标值','参照值','提升度%','参照对象'],['历史一致候选','21.43%\n12/56','21.43%\n12/56','0\n个百分点','自建固定温差阈值'],['审核反例正确拦截','2/2','0/2','不作泛化\n提升率','原两项反例'],['原始CSV至核查卡','1个完整\n任务通过','无导入\n入口','不适用','原预置案例流程']]
 while len(nested.rows)>len(vals):nested._tbl.remove(nested.rows[-1]._tr)
 for row,valsrow in zip(nested.rows,vals):
  for c,txt in zip(row.cells,valsrow):c.text=txt
 t.cell(18,2).add_paragraph('注：主口径56个基础日工况、14日期、2日期块；44个未决保留在分母。正常、泄漏召回仍为0；卡滞开度25%/75%分别1/14、11/14，不代表故障轻重。112条配对记录另作扰动回归，无新留出。2项反例和1项任务不构成泛化性能估计。')
 # All linked cells formatted once; do not create new nested styles or break merges.
 seen=set()
 for row in t.rows:
  for c in row.cells:
   if c._tc in seen:continue
   seen.add(c._tc)
   for p in c.paragraphs:pformat(p,10.5)
 for row in nested.rows:
  row.height=None
  for c in row.cells:
   for p in c.paragraphs:pformat(p,9)
 # Keep critical paired heading/body rows intact, permit long merged outer cell across pages.
 for i,row in enumerate(t.rows):
  if i not in [18]:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i in [5,7,9,11]:
   for p in row.cells[1].paragraphs:p.paragraph_format.keep_with_next=True
 sanitize_save(d,OUT/'附件1_能智核_项目简表.docx')

if __name__=="__main__":
 import runpy
 runpy.run_path(str(ROOT/"scripts/prepare_materials_round4.py"),run_name="__main__")
