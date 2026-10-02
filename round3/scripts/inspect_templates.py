from pathlib import Path
from docx import Document
from pypdf import PdfReader
import json
ROOT=Path(__file__).resolve().parents[1];out=[]
for p in sorted((ROOT/'materials/references').glob('*.docx')):
 d=Document(p);rows=[]
 for t in d.tables:
  rows.append([[c.text for c in r.cells] for r in t.rows])
 out.append({'file':p.name,'paragraphs':[x.text for x in d.paragraphs], 'tables':rows,
 'sections':[{'width':s.page_width.twips,'height':s.page_height.twips,'left':s.left_margin.twips,'right':s.right_margin.twips,'top':s.top_margin.twips,'bottom':s.bottom_margin.twips} for s in d.sections]})
(ROOT/'materials/working/template_slots.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
p=next((ROOT/'materials/references').glob('*.pdf'));pdf=PdfReader(p)
(ROOT/'materials/working/guide_text.txt').write_text('\n\n'.join(f'Page {i+1}\n'+(p.extract_text() or '') for i,p in enumerate(pdf.pages)),encoding='utf-8')
print(json.dumps(out,ensure_ascii=False,indent=2))
