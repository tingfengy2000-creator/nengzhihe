"""Template-derived anonymous competition drafts; no performance inventions."""
from pathlib import Path
import json,re,zipfile,io,hashlib
from lxml import etree
from docx import Document
from docx.shared import Pt,Inches,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1];REF=ROOT/'materials/references';OUT=ROOT/'materials/final';OUT.mkdir(exist_ok=True)
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main';NS={'w':W};counts={}
TITLE='能智核——公共建筑用能异常核验工作台'
def nchar(s):return len(re.sub(r'\s+','',s))
def font(run,size=12,bold=False,heading=False):
 run.font.name='SimHei' if heading else 'SimSun';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 rpr=run._element.get_or_add_rPr();rf=rpr.rFonts
 if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
 for k in list(rf.attrib):rf.attrib.pop(k)
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimHei' if heading else 'SimSun')
def pformat(p,size=12,heading=False):
 pf=p.paragraph_format;pf.space_before=Pt(5 if heading else 0);pf.space_after=Pt(6 if heading else 4);pf.line_spacing=1.22
 pf.keep_with_next=heading;pf.keep_together=heading;pf.widow_control=True;pf.first_line_indent=None
 for x in ['snapToGrid','pBdr','pageBreakBefore']:
  for el in p._p.xpath('w:pPr/w:'+x):el.getparent().remove(el)
 for r in p.runs:font(r,size,heading,heading)
def setup(i):
 d=Document(next(REF.glob(f'附件{i}*.docx')))
 if 'Title' not in d.styles:
  from docx.enum.style import WD_STYLE_TYPE
  d.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
 for style in d.styles:
  if style.type in (1,2):
   style.font.name='SimSun';style.font.size=Pt(12)
   if style.element.rPr is not None:
    for rf in style.element.rPr.findall(qn('w:rFonts')):
     for k in list(rf.attrib):rf.attrib.pop(k)
     for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimSun')
 for sec in d.sections:
  for grid in sec._sectPr.findall(qn('w:docGrid')):sec._sectPr.remove(grid)
  for area in [sec.header,sec.footer]:
   for p in area.paragraphs:pformat(p,9)
 d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.title=TITLE;d.core_properties.subject='智慧能源与环境';d.core_properties.comments=''
 return d
def para(d,text,size=12,heading=False):
 p=d.add_paragraph(text);pformat(p,size,heading);return p
def h(d,text):return para(d,text,14,True)
def sub(d,text):return para(d,text,12,True)
def newbody(i):
 d=setup(i)
 for el in list(d._element.body):
  if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
 para(d,f'附件{i}：',10.5)
 p=para(d,'第十二届中国研究生智慧城市技术与创意设计大赛',14,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,'创意设计赛项目'+('说明书' if i==2 else '商业计划书'),16,True);p.style=d.styles['Title'];pformat(p,16,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,TITLE,12,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p=para(d,'参赛方向：智慧能源与环境',10.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 return d
def page(d):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.add_run().add_break(__import__('docx').enum.text.WD_BREAK.PAGE)
def picture(d,path,caption):
 p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3);p.add_run().add_picture(str(path),width=Inches(5.70))
 p=para(d,caption,9.5);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(d,headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.autofit=False
 for j,x in enumerate(headers):t.rows[0].cells[j].text=x
 for row in rows:
  c=t.add_row().cells
  for j,x in enumerate(row):c[j].text=str(x)
 for i,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[j])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for tag in ['top','left','bottom','right']:
    line=OxmlElement('w:'+tag);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');line.set(qn('w:color'),'BBCBC5');b.append(line)
   tcpr.append(b)
   if i==0:s=OxmlElement('w:shd');s.set(qn('w:fill'),'EDF3EF');tcpr.append(s)
   for p in c.paragraphs:pformat(p,10.5,i==0)
 return t
def sanitize_save(d,path):
 buf=io.BytesIO();d.save(buf);zin=zipfile.ZipFile(buf);entries={n:zin.read(n) for n in zin.namelist() if not n.startswith('word/fonts/')}
 fontxml=etree.Element(qn('w:fonts'),nsmap={'w':W})
 for name in ['SimSun','SimHei']:
  f=etree.SubElement(fontxml,qn('w:font'));f.set(qn('w:name'),name)
 entries['word/fontTable.xml']=etree.tostring(fontxml,xml_declaration=True,encoding='UTF-8',standalone=True)
 entries.pop('word/_rels/fontTable.xml.rels',None)
 entries.pop('docProps/custom.xml',None)
 for name,raw in list(entries.items()):
  if not name.endswith(('.xml','.rels')):continue
  root=etree.fromstring(raw)
  if name=='[Content_Types].xml':
   for x in list(root):
    if x.get('Extension')=='odttf' or x.get('PartName','').startswith('/word/fonts/') or x.get('PartName')=='/docProps/custom.xml':root.remove(x)
  if name=='_rels/.rels':
   for x in list(root):
    if x.get('Type','').endswith('/custom-properties'):root.remove(x)
  if name=='docProps/app.xml':
   for x in list(root):
    local=etree.QName(x).localname
    if local=='Application':x.text='Office Open XML'
    elif local=='Pages':x.text=str(4 if '附件2' in path.name else 3)
    elif local in ['Words','Characters','CharactersWithSpaces','Lines','Paragraphs','Company','Manager','Template']:root.remove(x)
  # Strip subset fonts, theme fallback and unrequested template personal metadata.
  if name.startswith('word/'):
   for rf in root.findall('.//'+qn('w:rFonts')):
    ishead=rf.get(qn('w:ascii'))=='SimHei'
    for k in list(rf.attrib):rf.attrib.pop(k)
    for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'SimHei' if ishead else 'SimSun')
   for el in root.findall('.//'+qn('w:embedRegular'))+root.findall('.//'+qn('w:embedBold'))+root.findall('.//'+qn('w:embedItalic'))+root.findall('.//'+qn('w:embedBoldItalic')):el.getparent().remove(el)
  entries[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in entries.items():z.writestr(n,b)
def brief():
 d=setup(1);t=d.tables[0]
 for i,p in enumerate(d.paragraphs):
  if p.text.strip()=='(模板)':p.text=TITLE
  pformat(p,14 if i in [1,2] else 10.5,i in [1,2])
  if i in [1,2,3]:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for row in t.rows:
  row.height=None
  for pr in row._tr.xpath('w:trPr/w:trHeight'):pr.getparent().remove(pr)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 def cell(r,col,text):t.cell(r,col).text=text
 cell(0,1,TITLE)
 tracks=['低空经济与空天科','智慧交通与数字孪生','智能建造与城市更新','未来城市与太空人居','智慧文旅与乡村振兴','智慧能源与环境','韧性城市与智慧应急','AI与具身智能','智慧生活与大健康','城市运维与安全','智能材料与绿色建造','自由赛道','企业揭榜赛道','地方特色赛道']
 options=[f'{i+1}. {s}'+(' ☑' if i==5 else ' □') for i,s in enumerate(tracks)]
 cell(2,1,'1—11为定向赛道：\n'+'\n'.join('    '.join(options[i:i+2]) for i in range(0,14,2)))
 texts={6:'公共建筑用能告警发生后，管理人员仍需核对数据是否可信、设备是否处于可检验工况，以及下一步应补哪个点位或时段。能智核面向这一核查环节，提供数据质量与设备证据分开呈现的工作台。当前以公开单风道空调机组仿真数据实现核验，以BDG2独立展示真实用能回放，服务于有边界、可复查的异常核查。',
 8:'把“告警之后先查什么”转化为可操作流程：先校验单位和真实时间间隔，再检查每项故障假设的适用条件；有支持证据时给出候选，缺点位、缺工况或证据冲突时给出具体补证要求。已有本地软件、三个案例、公开原始CSV导入和可阅读核查卡。特色在证据约束的产品流程，不主张已证明新的诊断算法。',
 10:'用户选择LBNL单风道机组公开仿真日片段，映射时间列、点位、单位及采样间隔。程序显式转换温度与指令单位，将完整1分钟桶聚合为5分钟；缺失值保留，间隔不符则拒绝。服务、诊断证据、质量卡和导出复用同一质量对象。通过质量校验后，程序按启停、稳定运行、温差和关阀持续时长生成证据窗口，分别核验风路、盘管和运行工况。界面呈现支持与限制证据、分支状态及补证条件；修正数据或修改证据后真实重算。导出的HTML核查卡可离线阅读和打印，并记录源数据哈希、映射及运行编号。设备范围限于当前仿真配置，真实建筑回放不参与故障定位。',
 12:'潜在使用方为公共建筑运维团队及能源服务方，拟提供按设备点位配置的本地核查工具和接入服务。现阶段仅完成公开数据原型与功能验证，尚无试点、合作或收入事实。后续需经合法数据接入、工程师复核和现场验证后再确定收费。节能机会筛查为附属方向，当前不承诺节能量、人工工时或运维收益。'}
 for r,s in texts.items():cell(r,1,s)
 counts['简表']={str(r):{'chars':nchar(s),'cap':600 if r==10 else 300} for r,s in texts.items()}
 cell(13,2,'创意设计类；基础研究类；\n☑软硬件开发类；工程实施类')
 cell(14,2,'国际领先；国际先进；国内领先；国内先进；\n国内一般；☑无法判断')
 cell(15,2,'□硬件 ☑软件 □工艺 □方法 □服务 □商业模式\n□其他')
 cell(17,2,'国际产品/技术/标准；国内产品/技术/标准：未开展独立对标。\n自身修复前实现仅作工程对照，不作为行业参照。')
 cell(19,2,'☑质量提升：消除两个已复现的质量/时长错误。\n成本降低、效率提升、交付周期：未验证。\n☑新功能实现：原始CSV导入至可阅读核查卡。')
 nested=t.cell(18,2).tables[0]
 vals=[['指标名称','指标值','参照值','提升度%','参照对象'],['历史总体正确率','21.43%\n24/112','5.36%\n6/112','+16.07\n个百分点','自身修复前实现'],['审核反例正确拦截','2/2','0/2','不作泛化\n提升率','原两项反例'],['原始CSV至核查卡','1个完整\n任务通过','无导入\n入口','不适用','原预置案例流程']]
 while len(nested.rows)>len(vals):nested._tbl.remove(nested.rows[-1]._tr)
 for row,valsrow in zip(nested.rows,vals):
  for c,txt in zip(row.cells,valsrow):c.text=txt
 t.cell(18,2).add_paragraph('注：历史112条现仅作回归，含56个基础日工况及配对扰动；正常、泄漏召回仍为0，025/075基础工况正确数为1/14、11/14。自身工程修复不等于新算法。两个反例和一个导入任务不构成泛化性能估计。')
 # All linked cells formatted once; do not create new nested styles or break merges.
 seen=set()
 for row in t.rows:
  for c in row.cells:
   if c._tc in seen:continue
   seen.add(c._tc)
   for p in c.paragraphs:pformat(p,10.5)
 for row in nested.rows:
  row.height=None
  for c in row.cells:
   for p in c.paragraphs:pformat(p,9)
 # Keep critical paired heading/body rows intact, permit long merged outer cell across pages.
 for i,row in enumerate(t.rows):
  if i not in [18]:row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  if i in [5,7,9,11]:
   for p in row.cells[1].paragraphs:p.paragraph_format.keep_with_next=True
 sanitize_save(d,OUT/'附件1_能智核_项目简表.docx')

def technical():
 d=newbody(2);chapter=[]
 def txt(s):
  chapter.append(s);p=para(d,s,11);p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(4)
 def head(s):chapter.append(s);sub(d,s)
 def close(name,cap):counts[name]={'chars':sum(nchar(x) for x in chapter),'cap':cap};chapter.clear()
 h(d,'一、立项依据（不超过2000字）')
 txt('能智核面向公共建筑管理人员处理用能告警后的核查任务：先判断数据是否可用，再判断现有工况能否检验设备疑点，最后记录下一步动作。作品已完成限定公开数据范围的本地核验工作台，管理人员可将原始CSV映射为证据，导出可阅读、可追溯的核查卡。')
 txt('能耗变化本身不能直接指认设备故障。温度单位不一致、时间间隔错标或重复记录，会改变运行分钟和证据窗口；停机、持续制冷等工况又可能无法检验某项故障。若把这些问题混为一个“异常分数”，使用者难以判断应修数据、补时段还是检查设备。')
 txt('美国能源部EMIS资料将故障检测与诊断列为建筑能源信息管理的应用，说明其依赖多种运行数据，部分问题仍需额外分析或现场核查。能智核由此选择具体的告警核查环节，重点交付证据条件与行动清单。当前验证来自LBNL单风道机组公开仿真数据；BDG2仅承担真实用能曲线回放，两种数据不拼接成真实楼宇定位证据。')
 close('说明书一',2000)
 h(d,'二、项目创新内容（不超过3000字）')
 head('1．项目总体思路')
 txt('用户围绕一项疑点查看当前工况、数据质量、支持与限制证据及补证要求。三个演示分别展示真实用能回放、仿真工况核验和“修正数据后设备疑点仍存在”。新增原始CSV入口使这一流程可由用户操作重现。')
 picture(d,ROOT/'output/ui/workbench_top.png','图1 实际工作台：公开CSV完成映射后进入同一核验流程。')
 chapter.append('图1 实际工作台：公开CSV完成映射后进入同一核验流程。')
 head('2．可行性分析：项目的技术或实施可行性。')
 txt('系统由本地Python数值程序和浏览器界面组成，不需要付费API。已有17类允许点位的映射、单位转换、启停与窗口规则、参考包络及可追溯导出。公开源文件、字段说明、转换记录、反例和回归输入随复现包提供，数值均由程序计算。当前CSV入口仅接收LBNL单风道机组仿真日片段；接入其他楼宇须另行定义点位、校准与验证。')
 head('3．本项目的特色与创新之处。')
 txt('产品特色是把故障假设的可检验性作为结果的一部分：缺点位、工况不适用、证据冲突与已有支持证据分别解释，并写明应补的传感器或运行时段。数据质量与设备判断采用两个维度；同一质量对象贯穿服务、诊断、界面和导出，避免内部拒判而质量卡仍显示通过。正常参考包络内的观察只按证据强度表述，不能据此排除全部故障。')
 txt('这些特色属于核查流程设计与工程可靠性实现。现有证据尚未证明原创诊断算法、主动查询增益或人工运维提效；固定流程与自适应规则在历史主预算4下未显示核查收益。')
 close('说明书二',3000)
 h(d,'三、实施方案（不超过3000字）')
 head('输入与核验流程')
 picture(d,ROOT/'materials/working/flow.png','图2 与当前代码对应的输入、质量、适用条件、证据计算与导出流程。')
 chapter.append('图2 与当前代码对应的输入、质量、适用条件、证据计算与导出流程。')
 txt('导入时显式指定时间列、源点位、单位和原始采样间隔。公开样例保留源年度CSV中2018年4月10日的原始1440行、31列；该日期为已有开发演示，不计新验证。仅17类允许点位进入诊断，文件名、故障标签、故障设置和实际阀位反馈均不作输入。1分钟数据只在完整5点桶内取均值；任一点位缺值则保留该桶缺值。源数据时间不递增、重复、实际间隔与声明不符或映射越界时明确拒绝。')
 txt('质量有效后，程序剔除启动前30分钟，按真实时间连续性生成30分钟窗口，累计至少60分钟合格证据。风路核验要求室外与回风温差至少3℃并有对应指令参考；盘管核验要求风机稳定、零关阀指令持续至少75分钟，扣除15分钟等待后形成两个窗口，并有参考覆盖。未出现合格工况时保留未决。修改可用证据或修正重复数据后，重新计算并生成新运行编号。')
 head('已完成的功能与可靠性验证')
 table(d,['验证任务','实际结果'],[['单位反例','室外温度仅改标degF：质量、设备、导出一致拒判。'],['时间轴反例','1分钟时间戳冒充5分钟：拒判，累计运行分钟及窗口为0。'],['原始CSV完整操作','本地选择文件、映射、1440→288行、核验及HTML核查卡导出通过。'],['修正数据继续核查','删除12条完全重复记录后质量恢复；风阀疑点仍保留。'],['回归检查','29项可靠性检查通过；历史112条预测与上轮逐条一致。']],[1.45,4.20])
 chapter.extend(['验证任务 实际结果','单位反例：室外温度仅改标degF，质量、设备、导出一致拒判。','时间轴反例：1分钟冒充5分钟拒判，时长为0。','原始CSV：本地选择、映射、1440→288行、核查卡。','12条重复记录修正后仍有疑点。29项可靠性检查，112条历史回归。'])
 head('诊断结果与有效范围')
 txt('历史自身工程修复对照为6/112→24/112，总体正确率5.36%→21.43%，增加16.07个百分点；未决计入分母，异常未定位不算三分类正确。已判24条中错误0条；覆盖率与正确率数值相同，不作为独立进步重复计项。旧新共同正确4条，新增20条、退回未决2条，并非逐例改善。112条由56个基础日工况及其重复记录扰动组成，来自14个日期、同一仿真系统。')
 table(d,['数据源工况设置','基础日工况','正确','类别召回'],[['正常','14','0','0%'],['风阀卡滞025','14','1','7.14%'],['风阀卡滞075','14','11','78.57%'],['盘管阀泄漏','14','0','0%']],[2.35,1.1,0.75,1.45])
 chapter.append('数据源工况设置 基础日工况 正确 类别召回 正常14 0 0% 风阀卡滞025 14 1 7.14% 风阀卡滞075 14 11 78.57% 盘管阀泄漏14 0 0%')
 txt('全部112条的盘管合格窗口为0：16条无稳定运行，86条稳定运行中无零关阀指令，6条未通过等待等条件，4条等待后仅5分钟。因此该批输入不能充分检验正常与泄漏；025中9个基础工况虽落入正常参考包络，也不能据此排除卡滞。上述112条现仅作开发与回归，不能作为新性能留出。已完成的原始CSV演示也不构成新独立准确率证据。')
 txt('可复现的软件交付包括本地工作台、三个演示、原始CSV日片段与来源哈希、映射及导出卡、两项反例修复前后结果和逐例回归记录。网页修改证据或预算调用同一程序重算；完整信息参照获得全部允许证据。查询次数仅为实验记账，现场补采尚未实测，软件耗时不等于人工工时。')
 close('说明书三',3000)
 h(d,'四、应用前景分析（不超过500字）')
 txt('能智核拟服务于公共建筑管理人员与能源服务团队的异常初筛和复核交接，将“数据是否可信、当前能否检验、接下来补什么”落实为可保存的核查卡。现阶段可用于公开数据教学、规则验证和方案演示；尚不支持任意楼宇数据即接即用。后续价值应通过合法接入真实点位、现场核对和预先冻结的新留出验证确认，再讨论节能机会及商业服务。没有真实部署、付费客户或节能收益证明，亦不以当前候选结论直接控制设备。')
 close('说明书四',500)
 sub(d,'公开资料来源')
 para(d,'LBNL FDD 数据与设备说明：fdddata.lbl.gov（SDAHU公开仿真配置）；BDG2：github.com/buds-lab/building-data-genome-project-2。',9)
 para(d,'美国能源部《EMIS Primer》：betterbuildingssolutioncenter.energy.gov/sites/default/files/attachments/EMIS_Primer_Organizational_Use.pdf。',9)
 sanitize_save(d,OUT/'附件2_能智核_项目说明书.docx')

def business():
 d=newbody(3);chapter=[]
 def txt(s):chapter.append(s);para(d,s)
 def head(s):chapter.append(s);sub(d,s)
 def close(name,cap):counts[name]={'chars':sum(nchar(x) for x in chapter),'cap':cap};chapter.clear()
 h(d,'一、项目方案概述（不超过200字）')
 txt('能智核面向公共建筑用能告警后的核查工作，以数据质量和设备证据双维度呈现结果，明确当前工况能否检验疑点及下一步补证要求。已有本地工作台、三个演示、公开原始CSV映射和可阅读核查卡。当前仅验证限定仿真数据配置，拟在真实接入与现场验证后提供本地核查工具和实施服务。')
 close('商业一',200)
 h(d,'二、项目团队（不超过200字）')
 txt('项目工作覆盖数据处理、数值核验、软件界面与实验复核。现有技术交付可供逐项运行与检查；建筑运行知识、现场点位确认及服务交付仍需在后续应用验证中补齐。本匿名方案不列成员身份、培养单位和指导教师，也不以未证实的团队资历作为可行性依据。')
 close('商业二',200)
 h(d,'三、项目产品（服务）化（不超过2000字）')
 head('1.项目产品（服务）特性')
 txt('拟议产品承担告警后的核查交接：使用者导入合规日片段，确认时间、点位和单位映射，查看数据质量、工况适用性、支持或限制证据，并导出可阅读核查卡。核查卡包含运行编号、源文件哈希、计算窗口及补证条件，有助于明确本次判断依据。当前不连接楼控执行器，不自动下发控制指令。')
 txt('当前可交付的是限定范围的软件原型：LBNL单风道机组公开仿真CSV完成输入至导出任务，BDG2另作真实用能回放。用户不能把任意楼宇CSV上传后即视为可靠诊断；新设备接入需做点位和单位确认、参考校准与独立验证。质量错误时明确拒判，工况不适用时保留未决。')
 head('2.产品（服务）化实施计划')
 txt('第一阶段以现有交付建立可复核样板，由使用者复跑原始CSV、修改证据并核对输出。第二阶段拟在获得数据授权后，为一台同类设备建立点位字典和有效运行时段，与工程师核对候选和未决原因。第三阶段须在未参与调试的数据上检验错误率、覆盖率及实际核查成本，再决定是否进入试点。各阶段以证据通过为条件，目前没有已签订试点或合作。')
 close('商业三',2000)
 h(d,'四、项目产品（服务）市场与竞争（不超过2000字）')
 head('1.市场概述')
 txt('潜在用户包括校园、办公楼等公共建筑的运维团队，以及需要解释告警的能源服务方。美国能源部EMIS资料说明，能源信息管理可包含用能分析、故障检测与诊断，应用仍可能需要额外分析或现场核查。该资料说明任务存在，不等于本项目已获得市场规模或客户需求验证。本项目尚未开展足以推断市场规模的访谈或采购调查。')
 head('2.竞争优势分析')
 txt('人工表格便于核对，但需要人员自行维护单位、时间窗和判断依据；楼控告警提供运行信息，其数据口径和具体诊断能力取决于系统配置；专业FDD产品是后续应开展真实对标的对象。能智核当前可展示的特色，是把可检验条件、数据质量及补证要求置于同一核查流程，并保留可阅读记录。尚无独立产品对比，不能认定行业领先或人工提效。')
 table(d,['现有证据','支持的表述','不能推出'],[['2项审核反例被拦截','已消除对应质量/时长错误','全场景可靠性'],['1个公开CSV完整任务','输入至核查卡流程可用','任意楼宇即接即用'],['历史24/112正确','限定仿真下部分风阀证据成立','三分类能力充分或普遍节能']],[1.8,1.9,1.95])
 chapter.append('现有证据 支持的表述 不能推出 2项反例拦截不是全场景可靠性 1个CSV任务不是任意楼宇即用 历史24/112不能证明三分类或节能')
 head('3.项目实施风险及应对措施')
 txt('最直接风险是不可检验工况被误当成正常。历史112条中盘管合格窗口全部为0，正常与泄漏召回仍为0；风阀025与075的基础工况正确数分别为1/14与11/14，能力分布明显不均。产品因此明确未决原因和补证条件，不用放宽阈值代替证据。其他风险包括点位口径差异、真实标签不足及现场工作量未知，拟通过单设备接入审查、人工复核和冻结验证逐步确认。')
 close('商业四',2000)
 h(d,'五、商业模式（不超过2000字）')
 head('1.项目产品（服务）的开发、生产（服务）策略')
 txt('拟采用本地部署、按已验证设备配置交付的方式。软件提供映射、核验和核查卡；实施服务负责数据字典、单位与采样审查及证据条件确认。交付验收应同时检查能判断的案例和必须拒判的案例，不能只演示有利结果。现阶段服务规格与现场成本尚未测得。')
 head('2.项目产品（服务）的营销策略')
 txt('拟先面向具有数据授权与工程复核条件的使用方展示公开样板，沟通哪些告警需要核查、可提供哪些点位及如何验收。任何真实试点均需事先确定范围、责任及数据权限；目前无客户名单、合同或营销转化数据。展示以可运行任务和证据边界为主，不以算法领先或节能比例吸引采购。')
 head('3.项目产品（服务）获利方式')
 txt('可探索按设备配置的一次性接入服务和本地软件维护费用。收费标准、付费意愿与续费条件须由真实验证后确定，当前没有定价、收入和利润证据。只有在核查质量与实际工作成本均有记录后，才能判断是否具备可持续服务价值。')
 head('4.（若创业）企业发展计划')
 txt('暂无可确认的创业主体或融资事实。若后续验证成立，先形成单设备配置的可交付说明和责任边界，再评估组织与商业安排；在此之前不以规划代替实施成果。')
 close('商业五',2000)
 h(d,'六、预期经济效益分析（不超过500字）')
 txt('项目期望减少依据不明的反复核查，但尚未开展人工任务计时、现场节能测量或付费服务，不能给出确定收益。历史自身修复对照为5.36%→21.43%，增加16.07个百分点，只说明限定输入上的工程变化；程序查询单位与毫秒耗时均不能折算成人工成本或节能量。后续应记录同一任务下的已判错误、覆盖、实际查询与现场补采成本，并在采取运维措施后另行验证节能收益。当前经济效益仅为待检验的应用假设。')
 close('商业六',500)
 para(d,'公开资料：美国能源部 EMIS Primer，betterbuildingssolutioncenter.energy.gov；LBNL 公开FDD数据，fdddata.lbl.gov。',9)
 sanitize_save(d,OUT/'附件3_能智核_商业计划书.docx')

if __name__=='__main__':
 brief();technical();business()
 for name,x in counts.items():
  values=x.values() if name=='简表' else [x]
  for item in values:assert item['chars']<=item['cap'],(name,item)
 (ROOT/'materials/working/word_limits.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps(counts,ensure_ascii=False,indent=2))
